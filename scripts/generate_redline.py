"""
修改痕迹文档生成器 - 生成带 OOXML 修订标记的 .docx 文件
用法: python3 generate_redline.py <original.docx> --modifications <modifications.json> --output <redlined.docx>

生成标准 Word 修订模式文档（w:del 删除线 + w:ins 下划线插入），可在 Word/WPS 中正确显示。
每处修改附带批注（w:comment），标注修改理由和规则依据。
"""
import sys
import os
import json
import copy
import re
from datetime import datetime


def _get_document_default_rpr(doc):
    """从文档正文段落采样默认字体格式（字体、字号、字间距等）。
    
    优先从样式名为 Normal/正文 的段落采样，其次取最长的正文段落。
    确保新增段落的字体与原文档一致。
    """
    from docx.oxml.ns import qn
    
    best_rpr = None
    best_len = 0
    
    for para in doc.paragraphs:
        style_name = para.style.name if para.style else ""
        text = para.text.strip()
        if len(text) < 10:
            continue
        style_bonus = 1000 if style_name in ("Normal", "正文", "Body Text") else 0
        for run in para.runs:
            if not run.text.strip():
                continue
            rpr = run._element.find(qn('w:rPr'))
            if rpr is not None:
                score = len(run.text) + style_bonus
                if score > best_len:
                    best_len = score
                    best_rpr = rpr
    
    if best_rpr is not None:
        return copy.deepcopy(best_rpr)
    
    try:
        styles = doc.styles
        if styles:
            normal_style = styles.find("Normal") or styles.find("正文")
            if normal_style and hasattr(normal_style, 'element'):
                rpr = normal_style.element.find(qn('w:rPr'))
                if rpr is not None:
                    return copy.deepcopy(rpr)
    except Exception:
        pass
    return None


def _get_document_default_ppr(doc):
    """从文档正文段落采样默认段落格式（缩进、行距等）。"""
    from docx.oxml.ns import qn
    
    best_ppr = None
    best_len = 0
    
    for para in doc.paragraphs:
        style_name = para.style.name if para.style else ""
        text = para.text.strip()
        if len(text) < 10:
            continue
        style_bonus = 1000 if style_name in ("Normal", "正文", "Body Text") else 0
        ppr = para._element.find(qn('w:pPr'))
        if ppr is not None:
            score = len(text) + style_bonus
            if score > best_len:
                best_len = score
                best_ppr = ppr
    
    if best_ppr is not None:
        return copy.deepcopy(best_ppr)
    return None


def generate_redline(original_path, modifications_path, output_path):
    """生成修改痕迹文档"""
    from docx import Document
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    from lxml import etree

    doc = Document(original_path)

    with open(modifications_path, 'r', encoding='utf-8') as f:
        mods = json.load(f)

    # 兼容两种格式：列表直接传入 或 {"modifications": [...]}
    if isinstance(mods, list):
        modifications = mods
    else:
        modifications = mods.get("modifications", [])

    if not modifications:
        doc.save(output_path)
        return {"status": "success", "output": output_path, "modification_count": 0}

    # ===== 0. 采样文档默认格式（用于新增段落） =====
    default_rpr = _get_document_default_rpr(doc)
    default_ppr = _get_document_default_ppr(doc)

    if not modifications:
        doc.save(output_path)
        return {"status": "success", "output": output_path, "modification_count": 0}

    # ===== 1. 扫描整个文档，找到所有已有的最大 ID =====
    max_id = _find_max_id(doc)

    # ===== 2. 解析已有的 comments.xml（如果存在）=====
    existing_comments = _load_existing_comments(doc)

    timestamp = datetime.now().strftime("%Y-%m-%dT%H:%M:%SZ")
    author = "CSDN\u5408\u540c\u5ba1\u67e5"  # CSDN合同审查 (简体中文)

    applied_count = 0
    skipped = []
    comments_list = []

    for mod_idx, mod in enumerate(modifications):
        original_text = mod.get("original_text", "").strip()
        modified_text = mod.get("modified_text", "").strip()

        if not original_text:
            continue

        # 查找包含原始文本的段落
        target_para = None
        for para in doc.paragraphs:
            if original_text in para.text:
                target_para = para
                break

        if target_para is None:
            skipped.append({"text": original_text[:40], "reason": "未在段落中找到匹配文本"})
            continue

        # 检查目标文本是否在已删除的内容中（w:delText）
        p_elem = target_para._element
        if _is_in_deleted_content(p_elem, original_text):
            skipped.append({"text": original_text[:40], "reason": "文本已在先前的修订中被删除"})
            continue

        # ===== 查找包含目标文本的 run，支持跨 run 匹配 =====
        runs = p_elem.findall(qn('w:r'))
        matched_runs_info = _find_text_in_runs(runs, original_text, qn)

        if matched_runs_info is None:
            skipped.append({"text": original_text[:40], "reason": "文本在跨 run 或已删除内容中，无法精确匹配"})
            continue

        applied_count += 1

        # ===== 生成唯一的修订 ID =====
        max_id += 1
        del_id = str(max_id)
        max_id += 1
        ins_id = str(max_id)
        max_id += 1
        comment_id = str(max_id)

        # ===== 从原始 run 复制格式（保持字体一致）=====
        # 优先使用匹配到的run格式，如果为空则使用文档默认格式
        source_run = matched_runs_info['first_run']
        original_rpr = source_run.find(qn('w:rPr'))
        if original_rpr is not None:
            original_rpr = copy.deepcopy(original_rpr)
        elif default_rpr is not None:
            original_rpr = copy.deepcopy(default_rpr)
        else:
            original_rpr = OxmlElement('w:rPr')

        # ===== 创建修订标记 =====
        reason = mod.get("reason", "")
        rule_reference = mod.get("rule_reference", "")
        severity = mod.get("severity", "medium")
        category = mod.get("category", "")

        # 创建 w:del（删除原文）— 复制原始格式
        del_elem = OxmlElement('w:del')
        del_elem.set(qn('w:id'), del_id)
        del_elem.set(qn('w:author'), author)
        del_elem.set(qn('w:date'), timestamp)

        del_run = OxmlElement('w:r')
        if original_rpr is not None:
            del_rpr = copy.deepcopy(original_rpr)
        elif default_rpr is not None:
            del_rpr = copy.deepcopy(default_rpr)
        else:
            del_rpr = OxmlElement('w:rPr')
        del_run.append(del_rpr)
        del_text_elem = OxmlElement('w:delText')
        del_text_elem.set(qn('xml:space'), 'preserve')
        del_text_elem.text = original_text
        del_run.append(del_text_elem)
        del_elem.append(del_run)

        # ===== 检测是否为多段落插入（modified_text 含 \n）=====
        # 多段落时：第一段留在当前段落，后续段落创建新 w:p
        is_multi_paragraph = '\n' in modified_text

        if is_multi_paragraph:
            # 多段落插入：第一段的 ins
            segments = [s for s in modified_text.split('\n') if s.strip()]
            first_segment = segments[0]

            ins_elem = OxmlElement('w:ins')
            ins_elem.set(qn('w:id'), ins_id)
            ins_elem.set(qn('w:author'), author)
            ins_elem.set(qn('w:date'), timestamp)

            ins_run = OxmlElement('w:r')
            if original_rpr is not None:
                ins_rpr = copy.deepcopy(original_rpr)
            else:
                ins_rpr = OxmlElement('w:rPr')
            ins_run.append(ins_rpr)
            ins_text_elem = OxmlElement('w:t')
            ins_text_elem.set(qn('xml:space'), 'preserve')
            ins_text_elem.text = first_segment
            ins_run.append(ins_text_elem)
            ins_elem.append(ins_run)

            # 为后续段落创建独立的 w:p + w:ins
            extra_paragraphs = []
            # 判断段落样式：匹配中文序号标题（如 "五、知识产权"、"六、协议期限"）
            heading_pattern = re.compile(r'^[一二三四五六七八九十百]+、')
            # 获取原始段落的 pPr 作为标题段落参考
            source_ppr = p_elem.find(qn('w:pPr'))

            for seg in segments[1:]:
                is_heading = bool(heading_pattern.match(seg.strip()))

                # 创建新段落
                new_p = OxmlElement('w:p')

                # 设置段落属性：优先使用文档默认段落格式
                new_ppr = OxmlElement('w:pPr')
                if is_heading and source_ppr is not None:
                    # 标题段落：复制原始标题段落的 pPr（保持加粗、无缩进）
                    new_ppr = copy.deepcopy(source_ppr)
                elif default_ppr is not None:
                    # 正文段落：使用文档采样的默认段落格式
                    new_ppr = copy.deepcopy(default_ppr)
                    # 清除标题样式引用（pStyle），确保是正文格式
                    pstyle = new_ppr.find(qn('w:pStyle'))
                    if pstyle is not None:
                        new_ppr.remove(pstyle)
                    # 移除加粗
                    for rpr in new_ppr.findall(qn('w:rPr')):
                        for b_tag in rpr.findall(qn('w:b')) + rpr.findall(qn('w:bCs')):
                            rpr.remove(b_tag)
                elif source_ppr is not None:
                    # 兜底：从原始段落复制
                    new_ppr = copy.deepcopy(source_ppr)
                    pstyle = new_ppr.find(qn('w:pStyle'))
                    if pstyle is not None:
                        new_ppr.remove(pstyle)
                    for rpr in new_ppr.findall(qn('w:rPr')):
                        for b_tag in rpr.findall(qn('w:b')) + rpr.findall(qn('w:bCs')):
                            rpr.remove(b_tag)
                else:
                    # 最终兜底：手动设置基本段落格式
                    ind_elem = OxmlElement('w:ind')
                    ind_elem.set(qn('w:firstLine'), '400')
                    new_ppr.append(ind_elem)
                new_p.append(new_ppr)

                # 创建 w:ins 包裹内容
                max_id += 1
                seg_ins_id = str(max_id)
                seg_ins = OxmlElement('w:ins')
                seg_ins.set(qn('w:id'), seg_ins_id)
                seg_ins.set(qn('w:author'), author)
                seg_ins.set(qn('w:date'), timestamp)

                seg_run = OxmlElement('w:r')
                seg_rpr = OxmlElement('w:rPr')
                if original_rpr is not None:
                    seg_rpr = copy.deepcopy(original_rpr)
                elif default_rpr is not None:
                    seg_rpr = copy.deepcopy(default_rpr)
                # 正文段落去掉加粗（标题才有加粗）
                if not is_heading:
                    b_elem = seg_rpr.find(qn('w:b'))
                    if b_elem is not None:
                        seg_rpr.remove(b_elem)
                    bCs_elem = seg_rpr.find(qn('w:bCs'))
                    if bCs_elem is not None:
                        seg_rpr.remove(bCs_elem)
                seg_run.append(seg_rpr)
                seg_t = OxmlElement('w:t')
                seg_t.set(qn('xml:space'), 'preserve')
                seg_t.text = seg
                seg_run.append(seg_t)
                seg_ins.append(seg_run)
                new_p.append(seg_ins)

                extra_paragraphs.append(new_p)
        else:
            # 单段落插入（原有逻辑）
            ins_elem = OxmlElement('w:ins')
            ins_elem.set(qn('w:id'), ins_id)
            ins_elem.set(qn('w:author'), author)
            ins_elem.set(qn('w:date'), timestamp)

            ins_run = OxmlElement('w:r')
            if original_rpr is not None:
                ins_rpr = copy.deepcopy(original_rpr)
            elif default_rpr is not None:
                ins_rpr = copy.deepcopy(default_rpr)
            else:
                ins_rpr = OxmlElement('w:rPr')
            ins_run.append(ins_rpr)
            ins_text_elem = OxmlElement('w:t')
            ins_text_elem.set(qn('xml:space'), 'preserve')
            ins_text_elem.text = modified_text
            ins_run.append(ins_text_elem)
            ins_elem.append(ins_run)

        # ===== 创建批注（统一格式：改动操作 + 改动原因 + 改动依据）=====
        source_tag = mod.get("source", "")
        if not source_tag:
            source_tag = "审核要点" if rule_reference else "待法务确认"

        # 改动操作（简洁描述）
        if modified_text and modified_text != original_text:
            operation = "将原文修改为："
        else:
            operation = "删除原文。"

        # 清理改动原因中的规则编码（如 RULE-CONF-001），只保留可读文字
        clean_reason = re.sub(r'[、，；]?\s*RULE-[A-Z]+-\d+', '', reason)
        clean_reason = clean_reason.strip().rstrip('；').rstrip(';')

        # 风险等级标识
        severity_map = {"high": "⚠ 高风险", "medium": "⚡ 中风险", "low": "✎ 低风险"}
        severity_label = severity_map.get(severity, "✎ 低风险")

        # CUAD类别
        cuad_category = mod.get("cuad_category", "")
        cuad_label = f"[{cuad_category}]" if cuad_category else ""

        # 市场基准对比
        market_benchmark = mod.get("market_benchmark", "")
        negotiability = mod.get("negotiability", "")
        nego_map = {"high": "易谈判", "medium": "可协商", "low": "难改动", "none": "不可改"}
        nego_label = nego_map.get(negotiability, "可协商")

        # 构建批注（分层清晰，用分隔线区分）
        comment_text = f"【{severity_label}】{cuad_label} {category}\n"
        comment_text += f"━━━━━━━━━━━━━━━━━━\n"
        comment_text += f"{operation}\n"
        if modified_text:
            comment_text += f"{modified_text[:120]}"
            if len(modified_text) > 120:
                comment_text += "..."
            comment_text += "\n"
        comment_text += f"━━━━━━━━━━━━━━━━━━\n"
        # 原因：多个点（以数字编号分隔）每个点另起一行
        reason_lines = re.split(r'(?=\d+\.)', clean_reason)
        reason_lines = [l.strip() for l in reason_lines if l.strip()]
        comment_text += f"原因：\n"
        for line in reason_lines:
            comment_text += f"{line}\n"
        if market_benchmark:
            comment_text += f"━━━━━━━━━━━━━━━━━━\n"
            comment_text += f"市场基准：{market_benchmark}\n"
        comment_text += f"来源：{source_tag}"
        if negotiability:
            comment_text += f" | 可谈判性：{nego_label}"

        comments_list.append({
            "id": comment_id,
            "text": comment_text,
            "author": author,
            "date": timestamp
        })

        # ===== 替换：在第一个匹配 run 前插入批注起始，后面插入 del+ins+批注结束 =====
        first_run = matched_runs_info['first_run']
        crs = OxmlElement('w:commentRangeStart')
        crs.set(qn('w:id'), comment_id)
        first_run.addprevious(crs)

        last_run = matched_runs_info['last_run']
        last_run.addnext(del_elem)
        del_elem.addnext(ins_elem)

        # 批注范围结束
        cre = OxmlElement('w:commentRangeEnd')
        cre.set(qn('w:id'), comment_id)
        ins_elem.addnext(cre)

        # 批注引用 run（隐藏标记，WPS 用于关联批注）
        cr_run = OxmlElement('w:r')
        cr_rpr = OxmlElement('w:rPr')
        cr_run.append(cr_rpr)
        cr_ref = OxmlElement('w:commentReference')
        cr_ref.set(qn('w:id'), comment_id)
        cr_run.append(cr_ref)
        cre.addnext(cr_run)

        # 删除被替换的原始 runs
        for run in matched_runs_info['runs_to_remove']:
            p_elem.remove(run)

        # ===== 多段落插入：在当前段落后插入额外段落 =====
        if is_multi_paragraph:
            insert_after = p_elem
            for extra_p in extra_paragraphs:
                insert_after.addnext(extra_p)
                insert_after = extra_p

    # ===== 3. 合并并写入 comments =====
    all_comments = existing_comments + comments_list
    if all_comments:
        _write_comments(doc, all_comments, author, timestamp)

    # ===== 4. 启用修订模式 =====
    _enable_track_changes(doc)

    doc.save(output_path)

    return {
        "status": "success",
        "output": output_path,
        "modification_count": applied_count,
        "total_requested": len(modifications),
        "comments_added": len(comments_list),
        "skipped": skipped
    }


def _find_max_id(doc):
    """扫描文档中所有已有的修订 ID，返回最大值"""
    from docx.oxml.ns import qn
    from lxml import etree

    max_id = 0
    body = doc.element.body
    for elem in body.iter():
        id_val = elem.get(qn('w:id'))
        if id_val:
            try:
                max_id = max(max_id, int(id_val))
            except ValueError:
                pass

    try:
        for part in doc.part.package.iter_parts():
            if 'comments' in part.partname:
                try:
                    root = etree.fromstring(part.blob)
                    for elem in root.iter():
                        id_val = elem.get(qn('w:id'))
                        if id_val:
                            try:
                                max_id = max(max_id, int(id_val))
                            except ValueError:
                                pass
                except:
                    pass
    except:
        pass

    try:
        settings = doc.settings.element
        for elem in settings.iter():
            id_val = elem.get(qn('w:id'))
            if id_val:
                try:
                    max_id = max(max_id, int(id_val))
                except ValueError:
                    pass
    except:
        pass

    return max_id


def _load_existing_comments(doc):
    """加载已有的 comments.xml 内容"""
    from lxml import etree
    from docx.oxml.ns import qn

    existing = []
    try:
        for rel in doc.part.rels.values():
            if 'comments' in rel.reltype and 'commentsExtended' not in rel.reltype:
                try:
                    target_part = rel.target_part
                    root = etree.fromstring(target_part.blob)
                    W_NS = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
                    for comment_elem in root.findall(f'{{{W_NS}}}comment'):
                        comment_id = comment_elem.get(f'{{{W_NS}}}id')
                        texts = []
                        for t in comment_elem.iter(f'{{{W_NS}}}t'):
                            if t.text:
                                texts.append(t.text)
                        text = '\n'.join(texts)
                        if comment_id:
                            existing.append({
                                "id": comment_id,
                                "text": text,
                                "author": comment_elem.get(f'{{{W_NS}}}author', ''),
                                "date": comment_elem.get(f'{{{W_NS}}}date', '')
                            })
                except:
                    pass
    except:
        pass

    return existing


def _is_in_deleted_content(p_elem, text):
    """检查文本是否在 w:del 元素中（已被删除）"""
    from docx.oxml.ns import qn

    for del_elem in p_elem.findall(qn('w:del')):
        for dt in del_elem.iter(qn('w:delText')):
            if dt.text and text in dt.text:
                return True
    return False


def _is_in_del_element(run):
    """检查 run 是否在 w:del 元素内"""
    parent = run.getparent()
    while parent is not None:
        if parent.tag.endswith('}del') or parent.tag == 'del':
            return True
        parent = parent.getparent()
    return False


def _find_text_in_runs(runs, target_text, qn):
    """在多个 run 中查找目标文本（支持跨 run 匹配）"""
    normal_runs = []
    for run in runs:
        if _is_in_del_element(run):
            continue
        t_elems = run.findall(qn('w:t'))
        run_text = ''.join(t.text or '' for t in t_elems)
        normal_runs.append((run, run_text))

    full_text = ''.join(rt for _, rt in normal_runs)

    if target_text not in full_text:
        return None

    start_pos = full_text.index(target_text)
    end_pos = start_pos + len(target_text)

    current_pos = 0
    first_run = None
    last_run = None
    runs_to_remove = []

    for run, run_text in normal_runs:
        run_start = current_pos
        run_end = current_pos + len(run_text)

        if run_end <= start_pos:
            current_pos = run_end
            continue
        if run_start >= end_pos:
            break

        if first_run is None:
            first_run = run
        last_run = run
        runs_to_remove.append(run)
        current_pos = run_end

    if first_run is None:
        return None

    return {
        'first_run': first_run,
        'last_run': last_run,
        'runs_to_remove': runs_to_remove
    }


def _write_comments(doc, comments_list, author, timestamp):
    """写入批注到 docx 的 comments.xml，合并已有批注"""
    from docx.opc.part import Part
    from docx.opc.packuri import PackURI

    comments_xml = '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
    comments_xml += '<w:comments xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
    comments_xml += ' xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'

    for c in comments_list:
        text = c["text"].replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
        paragraphs = text.split("\n")
        para_xml = ""
        for p_text in paragraphs:
            para_xml += f'<w:p><w:r><w:t xml:space="preserve">{p_text}</w:t></w:r></w:p>'

        c_author = c.get("author", author)
        c_date = c.get("date", timestamp)
        comments_xml += f'<w:comment w:id="{c["id"]}" w:author="{c_author}" w:date="{c_date}">'
        comments_xml += para_xml
        comments_xml += '</w:comment>'

    comments_xml += '</w:comments>'

    existing_rel_id = None
    for rel_key, rel in list(doc.part.rels.items()):
        if 'comments' in rel.reltype and 'commentsExtended' not in rel.reltype:
            existing_rel_id = rel_key
            break

    if existing_rel_id:
        del doc.part.rels[existing_rel_id]

    comments_part_uri = PackURI('/word/comments.xml')
    content_type = 'application/vnd.openxmlformats-officedocument.wordprocessingml.comments+xml'

    comments_part = Part(
        comments_part_uri,
        content_type,
        comments_xml.encode('utf-8'),
        doc.part.package
    )

    doc.part.relate_to(comments_part,
                        'http://schemas.openxmlformats.org/officeDocument/2006/relationships/comments')


def _enable_track_changes(doc):
    """启用修订模式"""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    settings = doc.settings.element
    existing = settings.findall(qn('w:trackChanges'))
    if not existing:
        tc = OxmlElement('w:trackChanges')
        settings.append(tc)


if __name__ == "__main__":
    import argparse
    parser = argparse.ArgumentParser(description="修改痕迹文档生成器")
    parser.add_argument("original", help="原始合同文件路径 (.docx)")
    parser.add_argument("--modifications", required=True, help="修改建议 JSON 文件路径")
    parser.add_argument("--output", required=True, help="输出文件路径")
    args = parser.parse_args()

    result = generate_redline(args.original, args.modifications, args.output)
    print(json.dumps(result, ensure_ascii=False, indent=2))
